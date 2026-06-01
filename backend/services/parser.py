"""
Parser Module
Handles PDF and DOCX file parsing — extracts raw text.
No LLM used here. Pure file parsing.
"""
import io


def parse_resume(content: bytes, filename: str) -> str:
    """
    Parse resume file and return raw text.
    Supports: .pdf, .docx, .txt
    """
    if filename.endswith(".pdf"):
        return _parse_pdf(content)
    elif filename.endswith(".docx"):
        return _parse_docx(content)
    elif filename.endswith(".txt"):
        return content.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {filename}. Use PDF, DOCX, or TXT.")


def _parse_pdf(content: bytes) -> str:
    """Extract text from PDF using pdfplumber"""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        full_text = "\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("PDF appears to be image-based or empty. Please use a text-based PDF.")
        return full_text
    except ImportError:
        raise ImportError("pdfplumber not installed. Run: pip install pdfplumber")


def _parse_docx(content: bytes) -> str:
    """Extract text from DOCX using python-docx"""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip install python-docx")
